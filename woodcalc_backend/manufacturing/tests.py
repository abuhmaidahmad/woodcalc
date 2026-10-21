from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from rest_framework.test import APIClient

from tenants.models import Company, CompanyMembership
from .models import WorkOrder, WorkOrderItem


class WorkOrderBulkItemsTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='A', slug='a')
        user = get_user_model().objects.create_user(username='u', email='u@example.com', password='pw')
        CompanyMembership.objects.create(user=user, company=self.company, role=CompanyMembership.Role.OWNER)
        self.client = APIClient()
        self.client.force_authenticate(user)
        self.wo = WorkOrder.objects.create(tenant=self.company, order_number='WO-1', product_name='K')

    def test_creates_all_items_in_one_request(self):
        res = self.client.post(f'/api/manufacturing/work-orders/{self.wo.id}/items/', {'items': [
            {'description': 'Side 720x560x18mm', 'quantity': 4, 'unit': 'pcs'},
            {'description': 'Door 597x717x18mm', 'quantity': 2, 'unit': 'pcs'},
        ]}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(res.json(), {'created': 2})
        self.assertEqual(list(self.wo.items.order_by('id').values_list('description', 'quantity')),
                         [('Side 720x560x18mm', 4), ('Door 597x717x18mm', 2)])

    def test_rejects_another_tenants_work_order(self):
        other = Company.objects.create(name='B', slug='b')
        wo = WorkOrder.objects.create(tenant=other, order_number='WO-2', product_name='K')
        res = self.client.post(f'/api/manufacturing/work-orders/{wo.id}/items/', {'items': [
            {'description': 'x', 'quantity': 1, 'unit': 'pcs'},
        ]}, format='json')
        self.assertEqual(res.status_code, 404)
        self.assertFalse(WorkOrderItem.objects.filter(work_order=wo).exists())

    def test_single_item_endpoint_rejects_another_tenants_work_order(self):
        other = Company.objects.create(name='B', slug='b')
        wo = WorkOrder.objects.create(tenant=other, order_number='WO-2', product_name='K')
        res = self.client.post('/api/manufacturing/work-order-items/', {
            'work_order': wo.id, 'description': 'x', 'quantity': 1, 'unit': 'pcs',
        }, format='json')
        self.assertEqual(res.status_code, 400)
        self.assertFalse(WorkOrderItem.objects.filter(work_order=wo).exists())

    def test_single_item_endpoint_still_works_for_own_work_order(self):
        res = self.client.post('/api/manufacturing/work-order-items/', {
            'work_order': self.wo.id, 'description': 'x', 'quantity': 1, 'unit': 'pcs',
        }, format='json')
        self.assertEqual(res.status_code, 201, res.content)


class CuttingJobMaterialTests(TestCase):
    def setUp(self):
        from inventory.models import Material
        self.company = Company.objects.create(name='A', slug='a')
        user = get_user_model().objects.create_user(username='u', email='u@example.com', password='pw')
        CompanyMembership.objects.create(user=user, company=self.company, role=CompanyMembership.Role.OWNER)
        self.client = APIClient()
        self.client.force_authenticate(user)
        self.wo = WorkOrder.objects.create(tenant=self.company, order_number='WO-1', product_name='K')
        self.mat = Material.objects.create(tenant=self.company, sku='H1210_ST', name='Laminate',
                                           board_width=2800, board_height=2070, has_grain=True)

    def _job(self, material_id):
        return self.client.post('/api/manufacturing/cutting-jobs/', {
            'work_order': self.wo.id, 'material': material_id, 'thickness': '18', 'kerf': '4',
            'parts': [{'label': 'Door', 'width': '717', 'height': '447', 'quantity': 3, 'grain_locked': True}],
        }, format='json')

    def test_optimize_creates_stock_sheet_from_material_board(self):
        from .models import StockSheet
        job = self._job(self.mat.id)
        self.assertEqual(job.status_code, 201, job.content)
        res = self.client.post(f"/api/manufacturing/cutting-jobs/{job.json()['id']}/optimize/")
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(len(res.json()['layouts']), 1)
        sheet = StockSheet.objects.get(material=self.mat)
        self.assertEqual((sheet.width, sheet.height, sheet.thickness, sheet.grain_direction, sheet.tenant),
                         (2800, 2070, 18, 'length', self.company))

    def _door_placements(self):
        from .models import PartPlacement
        job = self._job(self.mat.id)
        res = self.client.post(f"/api/manufacturing/cutting-jobs/{job.json()['id']}/optimize/")
        self.assertEqual(res.status_code, 200, res.content)
        return list(PartPlacement.objects.filter(part__job_id=job.json()['id']))

    def test_grain_runs_along_board_long_side_when_entered_as_height(self):
        # 2070x2800: the long side (and the grain) is the board's height, so every
        # door's height (the part's width) must run along it.
        self.mat.board_width, self.mat.board_height = 2070, 2800
        self.mat.save()
        placements = self._door_placements()
        self.assertEqual(len(placements), 3)
        self.assertTrue(all(not p.rotated and (p.width, p.height) == (447, 717) for p in placements))

    def test_grained_material_overrides_stale_grainless_sheet(self):
        from .models import StockSheet
        StockSheet.objects.create(tenant=self.company, material=self.mat, thickness=18,
                                  width=2800, height=2070, grain_direction='none')
        from unittest import mock
        from . import cutting_service
        with mock.patch.object(cutting_service, 'pack_parts', wraps=cutting_service.pack_parts) as pack:
            self._door_placements()
        self.assertTrue(all(p.grain_locked for p in pack.call_args.args[0]))

    def test_material_without_grain_overrides_grained_sheet(self):
        from unittest import mock
        from .models import StockSheet
        from . import cutting_service
        self.mat.has_grain = False
        self.mat.save()
        StockSheet.objects.create(tenant=self.company, material=self.mat, thickness=18,
                                  width=2800, height=2070, grain_direction='length')
        with mock.patch.object(cutting_service, 'pack_parts', wraps=cutting_service.pack_parts) as pack:
            self._door_placements()
        self.assertFalse(any(p.grain_locked for p in pack.call_args.args[0]))

    def test_optimize_reports_material_without_board_size(self):
        self.mat.board_width = None
        self.mat.save()
        job = self._job(self.mat.id)
        res = self.client.post(f"/api/manufacturing/cutting-jobs/{job.json()['id']}/optimize/")
        self.assertEqual(res.status_code, 400)
        self.assertIn('H1210_ST', res.json()['detail'])

    def test_rejects_another_tenants_material(self):
        from inventory.models import Material
        other = Company.objects.create(name='B', slug='b')
        foreign = Material.objects.create(tenant=other, sku='X', name='X')
        res = self._job(foreign.id)
        self.assertEqual(res.status_code, 400)
        self.assertIn('material', res.json())


class CuttingJobStockTests(TestCase):
    """Optimizing a cutting job takes the sheets its layout uses out of stock."""

    def setUp(self):
        from inventory.models import Material
        self.company = Company.objects.create(name='A', slug='a')
        user = get_user_model().objects.create_user(username='u', email='u@example.com', password='pw')
        CompanyMembership.objects.create(user=user, company=self.company, role=CompanyMembership.Role.OWNER)
        self.client = APIClient()
        self.client.force_authenticate(user)
        self.wo = WorkOrder.objects.create(tenant=self.company, order_number='WO-1', product_name='K')
        self.hdf = Material.objects.create(tenant=self.company, sku='HDF3', name='HDF 3mm', unit='PCS',
                                           material_type='back', core_material='hdf',
                                           board_width=2440, board_height=1220, board_thickness=3,
                                           quantity_on_hand=10, reorder_level=5)

    def _job(self, qty, material=None):
        res = self.client.post('/api/manufacturing/cutting-jobs/', {
            'work_order': self.wo.id, 'material': (material or self.hdf).id, 'thickness': '3', 'kerf': '4',
            'parts': [{'label': 'Back', 'width': '1100', 'height': '1100', 'quantity': qty, 'grain_locked': False}],
        }, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        return res.json()['id']

    def _optimize(self, job_id):
        res = self.client.post(f'/api/manufacturing/cutting-jobs/{job_id}/optimize/')
        self.assertEqual(res.status_code, 200, res.content)
        return res.json()

    def _set_qty(self, job_id, qty):
        from .models import CuttingPart
        CuttingPart.objects.filter(job_id=job_id).update(quantity=qty)

    def test_optimize_deducts_sheets_and_logs_movement(self):
        from inventory.models import StockMovement
        data = self._optimize(self._job(4))  # 2 backs per 2440x1220 sheet -> 2 sheets
        self.assertEqual(len(data['layouts']), 2)
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, 8)
        mv = StockMovement.objects.get(material=self.hdf)
        self.assertEqual((mv.movement_type, mv.quantity), ('OUT', 2))
        self.assertEqual(data['stock']['used_by_job'], '2.00')
        self.assertEqual(data['stock']['quantity_on_hand'], '8.00')

    def test_reoptimize_books_only_the_difference(self):
        job_id = self._job(4)
        self._optimize(job_id)
        self._optimize(job_id)
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, 8)
        self._set_qty(job_id, 6)  # 3 sheets
        self._optimize(job_id)
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, 7)
        self._set_qty(job_id, 2)  # 1 sheet -> 2 come back
        self._optimize(job_id)
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, 9)

    def test_deleting_job_returns_stock(self):
        job_id = self._job(4)
        self._optimize(job_id)
        self.assertEqual(self.client.delete(f'/api/manufacturing/cutting-jobs/{job_id}/').status_code, 204)
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, 10)

    def test_dropping_below_reorder_level_raises_alert(self):
        self.assertFalse(self.hdf.alerts.filter(resolved=False).exists())
        data = self._optimize(self._job(12))  # 6 sheets -> 4 left, reorder at 5
        self.assertTrue(data['stock']['below_reorder_level'])
        self.assertTrue(self.hdf.alerts.filter(resolved=False).exists())

    def test_m2_stock_is_deducted_by_sheet_area(self):
        self.hdf.unit = 'M2'
        self.hdf.quantity_on_hand = 30
        self.hdf.save()
        self._optimize(self._job(4))  # 2 sheets x 2.9768 m2
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, Decimal('24.05'))

    def test_non_sheet_unit_is_left_alone(self):
        self.hdf.unit = 'KG'
        self.hdf.save()
        self._optimize(self._job(4))
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, 10)

    def test_new_run_for_same_work_order_replaces_earlier_job(self):
        # The optimizer page creates a new job on every run.
        self._optimize(self._job(4))   # 2 sheets
        self._optimize(self._job(6))   # 3 sheets, replaces the first run
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, 7)

    def test_other_work_orders_are_not_affected(self):
        self._optimize(self._job(4))
        other_wo = WorkOrder.objects.create(tenant=self.company, order_number='WO-2', product_name='K2')
        self.wo = other_wo
        self._optimize(self._job(4))
        self.hdf.refresh_from_db()
        self.assertEqual(self.hdf.quantity_on_hand, 6)


class CuttingJobXlsxExportTests(TestCase):
    """Several optimized jobs download as one workbook: summary + a tab per material."""

    def setUp(self):
        from inventory.models import Material
        self.company = Company.objects.create(name='A', slug='a')
        user = get_user_model().objects.create_user(username='u', email='u@example.com', password='pw')
        CompanyMembership.objects.create(user=user, company=self.company, role=CompanyMembership.Role.OWNER)
        self.client = APIClient()
        self.client.force_authenticate(user)
        self.wo = WorkOrder.objects.create(tenant=self.company, order_number='WO-1', product_name='K')
        self.ply = Material.objects.create(tenant=self.company, sku='wh.pw.18', name='White plywood',
                                           board_width=2440, board_height=1220, board_thickness=18)
        self.hdf = Material.objects.create(tenant=self.company, sku='HDF/3', name='HDF 3mm',
                                           board_width=2440, board_height=1220, board_thickness=3)

    def _optimized_job(self, material, thickness, parts):
        res = self.client.post('/api/manufacturing/cutting-jobs/', {
            'work_order': self.wo.id, 'material': material.id, 'thickness': thickness, 'kerf': '4', 'parts': parts,
        }, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        job_id = res.json()['id']
        self.assertEqual(self.client.post(f'/api/manufacturing/cutting-jobs/{job_id}/optimize/').status_code, 200)
        return job_id

    def _workbook(self, res):
        import io
        from openpyxl import load_workbook
        return load_workbook(io.BytesIO(b''.join(res.streaming_content) if res.streaming else res.content))

    def test_exports_summary_and_cut_list_per_material(self):
        ply = self._optimized_job(self.ply, '18', [
            {'label': 'Side panel', 'width': '720', 'height': '560', 'quantity': 4, 'grain_locked': False}])
        hdf = self._optimized_job(self.hdf, '3', [
            {'label': 'Back', 'width': '1100', 'height': '1100', 'quantity': 4, 'grain_locked': False}])
        res = self.client.get(f'/api/manufacturing/cutting-jobs/export-xlsx/?ids={ply},{hdf}')
        self.assertEqual(res.status_code, 200, res.content)
        self.assertIn('cutting_list_WO-1.xlsx', res['Content-Disposition'])
        wb = self._workbook(res)
        self.assertEqual(wb.sheetnames, ['Summary', 'wh.pw.18', 'HDF-3', 'Layouts'])
        rows = list(wb['Summary'].iter_rows(min_row=2, values_only=True))
        self.assertEqual(rows[0][:3], ('wh.pw.18', 'White plywood', 18))
        self.assertEqual(rows[0][4:], (4, 1))
        self.assertEqual(rows[1][4:], (4, 2))
        self.assertEqual(rows[2], ('Total', None, None, None, None, 3))
        cut = list(wb['wh.pw.18'].iter_rows(min_row=4, values_only=True))
        self.assertEqual(cut, [('Side panel', 720, 560, 4, 'No')])
        self.assertEqual(wb['Layouts'].max_row, 1 + 4 + 4)

    def test_csv_and_pdf_carry_arabic_part_names(self):
        job = self._optimized_job(self.ply, '18', [
            {'label': 'Side panel / جانب', 'width': '720', 'height': '560', 'quantity': 2, 'grain_locked': False}])
        res = self.client.get(f'/api/manufacturing/cutting-jobs/{job}/export-csv/')
        text = res.content.decode('utf-8')
        self.assertTrue(text.startswith('\ufeff'))  # so Excel reads it as UTF-8
        self.assertIn('Side panel / جانب', text)
        res = self.client.get(f'/api/manufacturing/cutting-jobs/{job}/export-pdf/')
        self.assertEqual(res.status_code, 200)
        self.assertIn(b'Amiri', res.content)

    def test_skips_another_tenants_jobs(self):
        from inventory.models import Material
        from .models import CuttingJob
        other = Company.objects.create(name='B', slug='b')
        wo = WorkOrder.objects.create(tenant=other, order_number='WO-9', product_name='K')
        mat = Material.objects.create(tenant=other, sku='X', name='X')
        job = CuttingJob.objects.create(work_order=wo, material=mat, thickness=18)
        res = self.client.get(f'/api/manufacturing/cutting-jobs/export-xlsx/?ids={job.id}')
        self.assertEqual(res.status_code, 404)


class PackPartsTests(SimpleTestCase):
    # Job 115's 50 pieces on a 2800x1220 board, 4mm kerf: ~4.2 boards of parts.
    JOB_115 = [(680, 1050, 2), (1077, 547, 2), (314, 300, 4), (350, 178, 2), (1077, 447, 7), (1019, 314, 2),
               (264, 250, 4), (230, 178, 2), (1080, 320, 2), (1055, 300, 4), (1800, 180, 2), (300, 900, 1),
               (1044, 264, 2), (1080, 250, 6), (1050, 250, 2), (1080, 180, 4), (1500, 180, 1), (1300, 180, 1)]

    def _parts(self, dims, grain_locked=False):
        from .cutting_service import Part
        return [Part(id=i, width=Decimal(w), height=Decimal(h), grain_locked=grain_locked)
                for i, (w, h, q) in enumerate(dims) for _ in range(q)]

    def _assert_valid(self, sheets, parts, sheet_w, sheet_h, kerf):
        placed = [pl for s in sheets for pl in s.placements]
        self.assertEqual(len(placed), len(parts))
        for s in sheets:
            pls = s.placements
            for a in pls:
                self.assertTrue(a.x >= 0 and a.y >= 0 and a.x + a.width <= sheet_w and a.y + a.height <= sheet_h)
                self.assertEqual((a.width, a.height),
                                 (a.part.height, a.part.width) if a.rotated else (a.part.width, a.part.height))
                if a.part.grain_locked:
                    self.assertFalse(a.rotated)
            for i, a in enumerate(pls):
                for b in pls[i + 1:]:
                    apart = (a.x + a.width + kerf <= b.x or b.x + b.width + kerf <= a.x or
                             a.y + a.height + kerf <= b.y or b.y + b.height + kerf <= a.y)
                    self.assertTrue(apart, f'{a} overlaps {b}')

    def test_job_115_fits_on_five_boards(self):
        from .cutting_service import pack_parts
        parts = self._parts(self.JOB_115)
        sheets, impossible = pack_parts(parts, Decimal(2800), Decimal(1220), Decimal(4))
        self.assertEqual(impossible, [])
        self.assertEqual(len(sheets), 5)
        self._assert_valid(sheets, parts, 2800, 1220, 4)

    def test_grain_locked_parts_are_never_turned(self):
        from .cutting_service import pack_parts
        parts = self._parts(self.JOB_115, grain_locked=True)
        sheets, impossible = pack_parts(parts, Decimal(2800), Decimal(1220), Decimal(4))
        self.assertEqual(impossible, [])
        self._assert_valid(sheets, parts, 2800, 1220, 4)

    def test_same_parts_give_the_same_layout(self):
        from .cutting_service import pack_parts
        def layout():
            sheets, _ = pack_parts(self._parts(self.JOB_115), Decimal(2800), Decimal(1220), Decimal(4))
            return [[(p.part.id, p.x, p.y, p.rotated) for p in s.placements] for s in sheets]
        self.assertEqual(layout(), layout())
